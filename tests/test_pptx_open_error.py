"""
Nightly debug run: a damaged or wrong .pptx upload (truncated download, a plain
zip or a renamed spreadsheet) showed the raw parser text to the visitor, e.g.
"Cannot open PowerPoint file: File is not a zip file" or "... file
'<_io.BytesIO object at 0x7f...>' is not a PowerPoint file, content type is ...",
always in English. It now shows a short text in the request's language; the raw
error goes to the server log only. The credit is still refunded.
"""
import io
import json
import zipfile

import pytest

import app as appmod


def _zip(files):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        for n, c in files.items():
            z.writestr(n, c)
    return b.getvalue()


_XLSX = _zip({
    "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                           '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>',
    "_rels/.rels": '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
    "xl/workbook.xml": "<w/>",
})

CASES = {
    "truncated": _zip({"a": "x" * 1000})[:60],
    "plain_zip": _zip({"notes.txt": "hello"}),
    "renamed_xlsx": _XLSX,
}


def _has_arabic(s):
    return any("؀" <= ch <= "ۿ" for ch in s or "")


def _errors(resp):
    out = []
    for line in resp.get_data(as_text=True).splitlines():
        if line.startswith("data:"):
            evt = json.loads(line[5:].strip())
            if "error" in evt:
                out.append(evt["error"])
    return out


@pytest.mark.parametrize("case", sorted(CASES))
@pytest.mark.parametrize("lang,arabic", [("ar", True), ("en", False)])
def test_broken_pptx_shows_friendly_localized_error(monkeypatch, case, lang, arabic):
    monkeypatch.setattr(appmod, "ollama_running", lambda: True)
    monkeypatch.setattr(appmod, "FREE_MODE", True)
    client = appmod.app.test_client()
    resp = client.post("/api/summarize-stream",
                       data={"file": (io.BytesIO(CASES[case]), "lecture.pptx"), "language": lang},
                       content_type="multipart/form-data",
                       headers={"X-Device-Id": "dev-pptx-open-%s-%s" % (case, lang)})
    assert resp.status_code == 200
    errs = _errors(resp)
    assert len(errs) == 1, errs
    msg = errs[0]
    for raw in ("Cannot open PowerPoint file", "zip file", "BytesIO", "0x", "Content_Types",
                "content type", "package"):
        assert raw not in msg, msg
    assert _has_arabic(msg) is arabic, msg


def test_valid_pptx_still_parses():
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Mitosis"
    s.placeholders[1].text = "Prophase, metaphase, anaphase, telophase"
    buf = io.BytesIO()
    prs.save(buf)
    slides = appmod.extract_slides(io.BytesIO(buf.getvalue()), filename="a.pptx")
    assert any("Mitosis" in (sl["title"] + sl["content"]) for sl in slides)
